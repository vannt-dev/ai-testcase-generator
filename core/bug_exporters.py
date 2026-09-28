"""
Tracker-neutral exports for bug reports: Markdown for pasting into
GitHub, GitLab, Azure DevOps or Jira Cloud, and Excel for sharing a set.
"""
import io
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED_FIELD_LABELS = {
    "title": "Title",
    "module": "Module",
    "steps_to_reproduce": "Steps to reproduce",
    "expected_result": "Expected result",
    "actual_result": "Actual result",
}

COLUMNS = [
    ("title", "Title", 40),
    ("module", "Module", 14),
    ("severity", "Severity", 11),
    ("priority", "Priority", 10),
    ("environment", "Environment", 22),
    ("preconditions", "Preconditions", 25),
    ("steps_to_reproduce", "Steps to Reproduce", 45),
    ("expected_result", "Expected Result", 30),
    ("actual_result", "Actual Result", 30),
    ("test_data", "Test Data", 20),
    ("related_test_id", "Related Test ID", 16),
    ("open_questions", "Open Questions", 35),
]

SEVERITY_COLORS = {
    "Critical": "FF9999",
    "Major": "FFC7CE",
    "Minor": "FFEB9C",
    "Trivial": "C6EFCE",
}

# Characters that start a Markdown block when they open a line: headings, quotes, lists,
# ~~~ fences, === setext headings, ___ rules and | tables.
_BLOCK_START = re.compile(r"^(\s*)([#>*+\-~=_|])", re.MULTILINE)
_ORDERED_START = re.compile(r"^(\s*\d+)([.)])", re.MULTILINE)


def _escape(text: str) -> str:
    # "<" would otherwise open an HTML tag that GitHub strips, e.g. "shows <null>".
    text = text.replace("\\", "\\\\").replace("`", "\\`").replace("<", "&lt;")
    text = _BLOCK_START.sub(r"\1\\\2", text)
    return _ORDERED_START.sub(r"\1\\\2", text)


def _steps(report: dict) -> list[str]:
    return [step.strip() for step in report.get("steps_to_reproduce", []) if step.strip()]


def file_stem(title: str) -> str:
    """ASCII download name from a title; Vietnamese diacritics are dropped, not mangled."""
    # NFKD splits most accented letters into base + combining mark; đ/Đ have no decomposition.
    ascii_title = unicodedata.normalize("NFKD", title.replace("đ", "d").replace("Đ", "D"))
    ascii_title = "".join(ch for ch in ascii_title if not unicodedata.combining(ch))
    slug = re.sub(r"[^A-Za-z0-9]+", "_", ascii_title).strip("_")[:50]
    return f"bug_{slug}" if slug else "bug_report"


def missing_required_fields(report: dict) -> list[str]:
    missing = []
    for key, label in REQUIRED_FIELD_LABELS.items():
        if key == "steps_to_reproduce":
            if not _steps(report):
                missing.append(label)
        elif not str(report.get(key, "")).strip():
            missing.append(label)
    return missing


def to_markdown(report: dict) -> str:
    title = " ".join(str(report.get("title", "")).split())
    meta = [
        f"**Severity:** {report.get('severity', '')}",
        f"**Priority:** {report.get('priority', '')}",
        f"**Module:** {_escape(report.get('module', ''))}",
    ]
    if report.get("environment", "").strip():
        meta.append(f"**Environment:** {_escape(report['environment'].strip())}")

    heading = _escape(title)
    # CommonMark drops a trailing run of "#" in a heading as its closing sequence.
    heading = re.sub(r"(?<!\\)(#+)$", lambda m: "\\#" * len(m.group(1)), heading)
    lines = [f"# {heading}", "", " · ".join(meta)]

    def section(heading: str, body: str) -> None:
        if body.strip():
            lines.extend(["", f"## {heading}", "", body.strip()])

    section("Preconditions", _escape(report.get("preconditions", "")))
    section(
        "Steps to Reproduce",
        "\n".join(f"{number}. {_escape(step)}" for number, step in enumerate(_steps(report), start=1)),
    )
    section("Expected Result", _escape(report.get("expected_result", "")))
    section("Actual Result", _escape(report.get("actual_result", "")))
    section("Test Data", _escape(report.get("test_data", "")))
    section("Related Test Case", _escape(report.get("related_test_id", "")))
    section(
        "Open Questions",
        "\n".join(f"- {_escape(q)}" for q in report.get("open_questions", []) if q.strip()),
    )
    return "\n".join(lines) + "\n"


def _cell_value(report: dict, key: str) -> str:
    if key == "steps_to_reproduce":
        return "\n".join(f"{number}. {step}" for number, step in enumerate(_steps(report), start=1))
    if key == "open_questions":
        return "\n".join(f"- {q}" for q in report.get("open_questions", []) if q.strip())
    return str(report.get(key, ""))


def to_excel(reports: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bug Reports"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col_idx, (_, header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, report in enumerate(reports, start=2):
        for col_idx, (key, _, _) in enumerate(COLUMNS, start=1):
            # openpyxl rejects control characters, e.g. ANSI colour codes in pasted terminal logs.
            text = ILLEGAL_CHARACTERS_RE.sub("", _cell_value(report, key))
            cell = ws.cell(row=row_idx, column=col_idx, value=text)
            # Report text is text even when it starts with '='.
            cell.data_type = "s"
            cell.alignment = wrap
            color = SEVERITY_COLORS.get(report.get("severity", "")) if key == "severity" else None
            if color:
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
