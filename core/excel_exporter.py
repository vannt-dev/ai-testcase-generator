"""
Exports a test case set (the dict returned by the AI) to an Excel (.xlsx)
file in a standard format, with basic styling for readability.
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from core.result_utils import covers_text
from core.traceability import build_traceability, has_traceability

COLUMNS = [
    ("test_id", "Test ID", 16),
    ("module", "Module", 14),
    ("title", "Test Title", 30),
    ("precondition", "Precondition", 25),
    ("steps", "Test Steps", 40),
    ("test_data", "Test Data", 20),
    ("expected_result", "Expected Result", 35),
    ("priority", "Priority", 10),
    ("type", "Type", 14),
    ("platform", "Platform", 12),
]

# Added to the Test Cases sheet, with a Traceability sheet, when the set has
# acceptance criteria or test cases that name the criteria they cover.
COVERS_COLUMN = ("covers", "Covers", 16)

UNCOVERED_FILL = "FFC7CE"

PRIORITY_COLORS = {
    "High": "FFC7CE",
    "Medium": "FFEB9C",
    "Low": "C6EFCE",
}


def export_to_excel(result: dict, sheet_name: str = "Test Cases") -> bytes:
    """
    Takes a dict {"test_cases": [...], "summary": {...}} and returns the
    .xlsx file content as bytes, ready to download or save to disk.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    traceable = has_traceability(result)
    columns = COLUMNS + [COVERS_COLUMN] if traceable else COLUMNS

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    # Header
    for col_idx, (_, header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data rows
    test_cases = result.get("test_cases", [])
    for row_idx, tc in enumerate(test_cases, start=2):
        for col_idx, (key, _, _) in enumerate(columns, start=1):
            value = covers_text(tc.get(key)) if key == "covers" else tc.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            # AI output and editor values are text even when they start with '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = wrap_alignment
            if key == "priority" and tc.get(key) in PRIORITY_COLORS:
                cell.fill = PatternFill(
                    start_color=PRIORITY_COLORS[tc[key]],
                    end_color=PRIORITY_COLORS[tc[key]],
                    fill_type="solid",
                )

    ws.freeze_panes = "A2"

    # Secondary sheet: Summary
    summary = result.get("summary", {})
    if summary:
        ws2 = wb.create_sheet("Summary")
        ws2.append(["Total test cases", summary.get("total", len(test_cases))])
        ws2.append([])
        ws2.append(["Type", "Count"])
        for k, v in summary.get("by_type", {}).items():
            ws2.append([k, v])
        ws2.append([])
        ws2.append(["Open questions to confirm with BA/Dev"])
        for q in summary.get("open_questions", []):
            ws2.append([q])
        for row in ws2.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        ws2.column_dimensions["A"].width = 45
        ws2.column_dimensions["B"].width = 15

    if traceable:
        _add_traceability_sheet(wb, result, header_font, header_fill, wrap_alignment)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _add_traceability_sheet(wb, result, header_font, header_fill, wrap_alignment) -> None:
    """One row per acceptance criterion with the test cases that cover it; gaps in red."""
    trace = build_traceability(result)
    ws = wb.create_sheet("Traceability")
    headers = [("Criterion", 14), ("Description", 50), ("Test IDs", 40), ("Test count", 11), ("Status", 28)]
    for col_idx, (header, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row in trace["criteria"]:
        if not row["test_ids"]:
            status = "No test case"
        elif row["listed"]:
            status = "Covered"
        else:
            status = "Covered (not in the criteria list)"
        ws.append([row["id"], row["text"], ", ".join(row["test_ids"]), len(row["test_ids"]), status])
        if not row["test_ids"]:
            fill = PatternFill(start_color=UNCOVERED_FILL, end_color=UNCOVERED_FILL, fill_type="solid")
            for cell in ws[ws.max_row]:
                cell.fill = fill

    if trace["unlinked_test_ids"]:
        ws.append([])
        ws.append(["Test cases not linked to a criterion"])
        for test_id in trace["unlinked_test_ids"]:
            ws.append([test_id])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = wrap_alignment
    ws.freeze_panes = "A2"
