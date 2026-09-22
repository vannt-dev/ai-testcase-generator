"""Parses uploaded test case files (.xlsx/.csv) into raw rows, independent of any target schema."""
import csv
import io
from collections import Counter

import openpyxl

MAX_IMPORTED_ROWS = 500
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMPORTED_COLUMNS = 256
MAX_FIELD_CHARS = 131072


class FileImportError(ValueError):
    """File import error with a user-friendly message for the UI."""


def parse_uploaded_file(uploaded_file) -> list[dict]:
    """
    uploaded_file: an object exposing `.name` (str) and `.getvalue()`
    (bytes) — matches Streamlit's UploadedFile.
    Returns a list of dicts keyed by the file's own column headers.
    Raises FileImportError on an unsupported extension, an empty or
    unreadable file, or more than MAX_IMPORTED_ROWS data rows.
    """
    name = getattr(uploaded_file, "name", "")
    suffix = name.rsplit(".", 1)[-1].lower() if "." in name else ""

    # Reject an unsupported extension before reading a single byte.
    if suffix not in ("csv", "xlsx"):
        raise FileImportError(
            f"Unsupported file type '.{suffix}'. Please upload a .xlsx or .csv file."
        )

    raw_bytes = uploaded_file.getvalue()
    if len(raw_bytes) > MAX_UPLOAD_BYTES:
        raise FileImportError("The uploaded file exceeds the 10 MiB size limit.")
    try:
        rows = _parse_csv(raw_bytes) if suffix == "csv" else _parse_xlsx(raw_bytes)
    except FileImportError:
        raise
    except (csv.Error, ValueError, OSError) as error:
        # Parsing is lazy: failures can occur while reading headers or later rows.
        raise FileImportError(f"Could not read the uploaded {suffix.upper()} file: {error}") from error

    if not rows:
        raise FileImportError("The uploaded file has no data rows.")
    return rows


def _too_many_rows_error() -> FileImportError:
    return FileImportError(
        f"The uploaded file has more than {MAX_IMPORTED_ROWS:,} rows, which exceeds the "
        f"{MAX_IMPORTED_ROWS:,} row limit. Please split it into smaller files."
    )


def _normalize_headers(raw_headers) -> list[str]:
    """Return stripped string headers and reject ambiguous column names."""
    _validate_cells(raw_headers)
    headers = [str(header).strip() if header is not None else "" for header in raw_headers]
    usable_headers = [header for header in headers if header]
    duplicates = sorted(
        header for header, count in Counter(usable_headers).items() if count > 1
    )
    if duplicates:
        raise FileImportError(
            "The uploaded file has duplicate column headers: " + ", ".join(duplicates)
        )
    if not usable_headers:
        raise FileImportError("The uploaded file has no named columns.")
    return headers


def _validate_cells(values) -> None:
    if len(values) > MAX_IMPORTED_COLUMNS:
        raise FileImportError(f"The uploaded file exceeds the {MAX_IMPORTED_COLUMNS} column limit.")
    if any(isinstance(value, str) and len(value) > MAX_FIELD_CHARS for value in values):
        raise FileImportError(f"An uploaded field exceeds the {MAX_FIELD_CHARS:,} character limit.")


def _parse_csv(raw_bytes: bytes) -> list[dict]:
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise FileImportError(f"Could not read the CSV file as UTF-8: {error}") from error
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        return []
    headers = _normalize_headers(reader.fieldnames)
    reader.fieldnames = headers
    rows = []
    for row in reader:
        if None in row:
            raise FileImportError(
                "A CSV data row has more values than the header row."
            )
        _validate_cells(list(row.values()))
        if len(rows) >= MAX_IMPORTED_ROWS:
            # Stop as soon as the cap is exceeded so a hostile file cannot
            # be fully materialized in memory first.
            raise _too_many_rows_error()
        rows.append(dict(row))
    return rows


def _parse_xlsx(raw_bytes: bytes) -> list[dict]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception as error:
        raise FileImportError(f"Could not read the Excel file: {error}") from error

    try:
        sheet = workbook.active
        if sheet is None:
            raise FileImportError("The Excel file has no active worksheet.")
        if sheet.max_column and sheet.max_column > MAX_IMPORTED_COLUMNS:
            raise FileImportError(f"The uploaded file exceeds the {MAX_IMPORTED_COLUMNS} column limit.")
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            headers = _normalize_headers(next(rows_iter))
        except StopIteration:
            return []

        rows = []
        for values in rows_iter:
            _validate_cells(values)
            if all(v is None for v in values):
                continue
            if len(rows) >= MAX_IMPORTED_ROWS:
                raise _too_many_rows_error()
            rows.append({h: (values[i] if i < len(values) else None) for i, h in enumerate(headers)})
        return rows
    except FileImportError:
        raise
    except Exception as error:
        # XML/ZIP errors can surface only when a read-only worksheet is consumed.
        raise FileImportError(f"Could not read the Excel file: {error}") from error
    finally:
        workbook.close()
