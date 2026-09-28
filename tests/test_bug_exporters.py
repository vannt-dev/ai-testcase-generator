import io

import openpyxl

from core.bug_exporters import file_stem, missing_required_fields, to_excel, to_markdown


def _report(**overrides):
    report = {
        "title": "Checkout freezes when paying with an expired card",
        "module": "Checkout",
        "severity": "Major",
        "priority": "High",
        "environment": "Android 15, app 3.2.0",
        "preconditions": "Logged in with one item in the cart",
        "steps_to_reproduce": ["Open the cart", "Tap Pay", "Enter an expired card"],
        "expected_result": "An 'expired card' error is shown",
        "actual_result": "The app freezes",
        "test_data": "Card expiry 01/20",
        "related_test_id": "TC_PAY_003",
        "open_questions": ["Does it happen on iOS?"],
    }
    report.update(overrides)
    return report


def _sheet(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def test_markdown_has_sections_in_order_with_numbered_steps():
    md = to_markdown(_report())

    assert md.startswith("# Checkout freezes when paying with an expired card\n")
    assert "**Severity:** Major · **Priority:** High · **Module:** Checkout" in md
    assert "**Environment:** Android 15, app 3.2.0" in md
    order = [
        "## Preconditions",
        "## Steps to Reproduce",
        "## Expected Result",
        "## Actual Result",
        "## Test Data",
        "## Related Test Case",
        "## Open Questions",
    ]
    positions = [md.index(heading) for heading in order]
    assert positions == sorted(positions)
    assert "1. Open the cart\n2. Tap Pay\n3. Enter an expired card" in md
    assert "- Does it happen on iOS?" in md


def test_markdown_leaves_out_empty_optional_sections():
    md = to_markdown(
        _report(environment="", preconditions="", test_data="", related_test_id="", open_questions=[])
    )

    for heading in ("## Preconditions", "## Test Data", "## Related Test Case", "## Open Questions"):
        assert heading not in md
    assert "**Environment:**" not in md
    assert "## Expected Result" in md


def test_markdown_escapes_user_markdown():
    md = to_markdown(
        _report(
            title="Error `null`\nshown",
            actual_result="# Not a heading\n- not a list\n1. not a list",
            steps_to_reproduce=["> quote", "Use `code`"],
        )
    )

    assert md.startswith("# Error \\`null\\` shown\n")
    assert "\\# Not a heading\n\\- not a list\n1\\. not a list" in md
    assert "1. \\> quote\n2. Use \\`code\\`" in md


def test_missing_required_fields_names_blank_fields_in_display_order():
    report = _report(title=" ", expected_result="", steps_to_reproduce=["", "  "])

    assert missing_required_fields(report) == ["Title", "Steps to reproduce", "Expected result"]
    assert missing_required_fields(_report()) == []


def test_excel_has_headers_one_row_per_report_and_severity_fill():
    sheet = _sheet(to_excel([_report(), _report(title="Second", severity="Trivial")]))

    headers = [cell.value for cell in sheet[1]]
    assert headers[:4] == ["Title", "Module", "Severity", "Priority"]
    assert "Steps to Reproduce" in headers
    assert sheet.max_row == 3
    steps_col = headers.index("Steps to Reproduce") + 1
    assert sheet.cell(row=2, column=steps_col).value == "1. Open the cart\n2. Tap Pay\n3. Enter an expired card"
    severity_col = headers.index("Severity") + 1
    assert sheet.cell(row=2, column=severity_col).fill.start_color.rgb.endswith("FFC7CE")
    assert sheet.cell(row=3, column=severity_col).fill.start_color.rgb.endswith("C6EFCE")


def test_excel_keeps_formula_like_text_literal():
    sheet = _sheet(to_excel([_report(actual_result="=SUM(A1:A2)", test_data="=1+1")]))

    values = [cell.value for cell in sheet[2]]
    assert "=SUM(A1:A2)" in values
    assert all(cell.data_type == "s" for cell in sheet[2] if isinstance(cell.value, str))


def test_markdown_and_excel_keep_non_ascii_text():
    report = _report(title="Ứng dụng bị treo khi thanh toán", actual_result="Không hiện thông báo lỗi")

    assert "# Ứng dụng bị treo khi thanh toán" in to_markdown(report)
    sheet = _sheet(to_excel([report]))
    assert sheet.cell(row=2, column=1).value == "Ứng dụng bị treo khi thanh toán"


def test_excel_strips_control_characters_from_pasted_logs():
    sheet = _sheet(to_excel([_report(actual_result="log \x1b[31mERR\x1b[0m\x00")]))

    headers = [cell.value for cell in sheet[1]]
    actual = sheet.cell(row=2, column=headers.index("Actual Result") + 1).value
    assert actual == "log [31mERR[0m"


def test_markdown_escapes_fences_setext_rules_tables_and_tags():
    md = to_markdown(
        _report(actual_result="~~~\ntext\n===\n___\n| a | b |\nshows <null>")
    )

    assert "\~~~\ntext\n\===\n\___\n\| a | b |\nshows &lt;null>" in md


def test_markdown_keeps_a_trailing_hash_in_the_title():
    md = to_markdown(_report(title="Cart badge shows #"))

    assert md.startswith("# Cart badge shows \#\n")


def test_file_stem_transliterates_vietnamese_titles():
    assert file_stem("Ứng dụng bị treo khi thanh toán") == "bug_Ung_dung_bi_treo_khi_thanh_toan"
    assert file_stem("Đăng nhập thất bại") == "bug_Dang_nhap_that_bai"
    assert file_stem("!!!") == "bug_report"


def test_markdown_does_not_double_escape_a_hash_only_title():
    assert to_markdown(_report(title="#")).startswith("# \#\n")
