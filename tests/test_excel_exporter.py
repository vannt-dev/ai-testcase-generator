import copy
import io

import pytest

from openpyxl import load_workbook

from core.excel_exporter import export_to_excel


@pytest.mark.parametrize("text", ["=1+1", "+1", "-1", "@test", "=HYPERLINK(\"https://example.invalid\")"])
def test_export_keeps_generated_edited_and_summary_text_literal(text):
    result = copy.deepcopy(SAMPLE_RESULT)
    for field in result["test_cases"][0]:
        result["test_cases"][0][field] = text
    result["summary"]["open_questions"] = [text]
    result["summary"]["by_type"] = {text: 1}
    workbook = load_workbook(io.BytesIO(export_to_excel(result)))
    for cell in workbook["Test Cases"][2]:
        assert cell.value == text
        assert cell.data_type == "s"
    summary_cells = [cell for row in workbook["Summary"] for cell in row if cell.value == text]
    assert len(summary_cells) == 2
    assert all(cell.data_type == "s" for cell in summary_cells)
    assert workbook["Summary"]["B1"].data_type == "n"


SAMPLE_RESULT = {
    "test_cases": [
        {
            "test_id": "TC_LOGIN_001",
            "module": "Login",
            "title": "Successful login with valid phone number + OTP",
            "precondition": "Account is registered, not logged in",
            "steps": "1. Enter phone number\n2. Enter correct OTP\n3. Tap Login",
            "test_data": "Phone: 0912345678, OTP: 123456",
            "expected_result": "Login succeeds, redirected to the Home screen",
            "priority": "High",
            "type": "Positive",
            "platform": "All",
        },
        {
            "test_id": "TC_LOGIN_002",
            "module": "Login",
            "title": "Login fails with an incorrect OTP",
            "precondition": "Account is registered",
            "steps": "1. Enter phone number\n2. Enter incorrect OTP",
            "test_data": "OTP: 000000",
            "expected_result": "Shows the error 'Incorrect OTP'",
            "priority": "Medium",
            "type": "Negative",
            "platform": "All",
        },
    ],
    "summary": {
        "total": 2,
        "by_type": {"positive": 1, "negative": 1},
        "open_questions": ["How long until the OTP expires?"],
    },
}


def test_export_to_excel_returns_valid_xlsx_bytes():
    excel_bytes = export_to_excel(SAMPLE_RESULT)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = load_workbook(io.BytesIO(excel_bytes))
    assert "Test Cases" in wb.sheetnames
    assert "Summary" in wb.sheetnames


def test_test_cases_sheet_has_correct_header_and_rows():
    excel_bytes = export_to_excel(SAMPLE_RESULT)
    ws = load_workbook(io.BytesIO(excel_bytes))["Test Cases"]

    header = [cell.value for cell in ws[1]]
    assert header == [
        "Test ID", "Module", "Test Title", "Precondition", "Test Steps",
        "Test Data", "Expected Result", "Priority", "Type", "Platform",
    ]

    assert ws.cell(row=2, column=1).value == "TC_LOGIN_001"
    assert ws.cell(row=2, column=8).value == "High"
    assert ws.cell(row=3, column=1).value == "TC_LOGIN_002"


def test_priority_cells_are_color_coded():
    excel_bytes = export_to_excel(SAMPLE_RESULT)
    ws = load_workbook(io.BytesIO(excel_bytes))["Test Cases"]

    high_priority_cell = ws.cell(row=2, column=8)  # priority = High
    assert high_priority_cell.fill.start_color.rgb == "00FFC7CE"

    medium_priority_cell = ws.cell(row=3, column=8)  # priority = Medium
    assert medium_priority_cell.fill.start_color.rgb == "00FFEB9C"


def test_summary_sheet_contains_totals_and_open_questions():
    excel_bytes = export_to_excel(SAMPLE_RESULT)
    ws = load_workbook(io.BytesIO(excel_bytes))["Summary"]

    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    flat_values = [v for row in rows for v in row if v is not None]

    assert 2 in flat_values  # total
    assert "How long until the OTP expires?" in flat_values


def test_export_with_no_test_cases_still_produces_file():
    result = {"test_cases": [], "summary": {"total": 0, "by_type": {}, "open_questions": []}}
    excel_bytes = export_to_excel(result)
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb["Test Cases"]
    assert ws.max_row == 1  # header row only
