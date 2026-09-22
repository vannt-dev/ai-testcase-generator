import csv
import io
from unittest.mock import Mock
from xml.etree.ElementTree import ParseError

import openpyxl
import pytest

from core.file_import import FileImportError, MAX_IMPORTED_ROWS, parse_uploaded_file


class FakeUploadedFile:
    def __init__(self, name: str, content: bytes):
        self.name = name
        self._content = content

    def getvalue(self) -> bytes:
        return self._content


def test_long_csv_field_is_reported_as_a_user_facing_import_error():
    content = ("Title\n" + "x" * (csv.field_size_limit() + 1)).encode()
    with pytest.raises(FileImportError, match="Could not read.*CSV"):
        parse_uploaded_file(FakeUploadedFile("long.csv", content))


@pytest.mark.parametrize("header_read", [True, False])
def test_lazy_excel_errors_are_reported_and_workbook_is_closed(monkeypatch, header_read):
    def rows():
        if not header_read:
            yield ("Title",)
        raise ParseError("malformed worksheet")

    workbook = Mock()
    workbook.active.max_column = 1
    workbook.active.iter_rows.return_value = rows()
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *args, **kwargs: workbook)
    with pytest.raises(FileImportError, match="Could not read the Excel file"):
        parse_uploaded_file(FakeUploadedFile("broken.xlsx", b"fixture"))
    workbook.close.assert_called_once()


def test_upload_size_limit_precedes_parsing():
    from core.file_import import MAX_UPLOAD_BYTES

    with pytest.raises(FileImportError, match="size limit"):
        parse_uploaded_file(FakeUploadedFile("large.csv", b"x" * (MAX_UPLOAD_BYTES + 1)))


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_upload_column_limit(extension):
    from core.file_import import MAX_IMPORTED_COLUMNS

    headers = [f"field{i}" for i in range(MAX_IMPORTED_COLUMNS + 1)]
    content = (",".join(headers) + "\n" + ",".join(headers)).encode() if extension == "csv" else _make_xlsx_bytes(headers, [headers])
    with pytest.raises(FileImportError, match="column limit"):
        parse_uploaded_file(FakeUploadedFile(f"wide.{extension}", content))


def _make_xlsx_bytes(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parse_uploaded_xlsx_returns_rows_keyed_by_header():
    content = _make_xlsx_bytes(
        ["ID", "Title"],
        [["TC_001", "Login works"], ["TC_002", "Login fails"]],
    )
    uploaded = FakeUploadedFile("cases.xlsx", content)

    rows = parse_uploaded_file(uploaded)

    assert rows == [
        {"ID": "TC_001", "Title": "Login works"},
        {"ID": "TC_002", "Title": "Login fails"},
    ]


def test_parse_uploaded_csv_returns_rows_keyed_by_header():
    content = "ID,Title\nTC_001,Login works\nTC_002,Login fails\n".encode("utf-8")
    uploaded = FakeUploadedFile("cases.csv", content)

    rows = parse_uploaded_file(uploaded)

    assert rows == [
        {"ID": "TC_001", "Title": "Login works"},
        {"ID": "TC_002", "Title": "Login fails"},
    ]


def test_parse_uploaded_file_rejects_unsupported_extension():
    uploaded = FakeUploadedFile("cases.txt", b"whatever")

    with pytest.raises(FileImportError, match="Unsupported file type"):
        parse_uploaded_file(uploaded)


def test_parse_uploaded_file_rejects_empty_file():
    content = _make_xlsx_bytes(["ID", "Title"], [])
    uploaded = FakeUploadedFile("empty.xlsx", content)

    with pytest.raises(FileImportError, match="no data rows"):
        parse_uploaded_file(uploaded)


def test_parse_uploaded_file_rejects_empty_csv():
    with pytest.raises(FileImportError, match="no data rows"):
        parse_uploaded_file(FakeUploadedFile("empty.csv", b""))


def test_parse_uploaded_file_rejects_too_many_rows():
    content = _make_xlsx_bytes(["ID"], [[f"TC_{i}"] for i in range(MAX_IMPORTED_ROWS + 1)])
    uploaded = FakeUploadedFile("huge.xlsx", content)

    with pytest.raises(FileImportError, match="row limit"):
        parse_uploaded_file(uploaded)


def test_parse_uploaded_file_accepts_exactly_max_rows():
    content = _make_xlsx_bytes(["ID"], [[f"TC_{i}"] for i in range(MAX_IMPORTED_ROWS)])
    uploaded = FakeUploadedFile("boundary.xlsx", content)

    rows = parse_uploaded_file(uploaded)

    assert len(rows) == MAX_IMPORTED_ROWS


def test_parse_uploaded_csv_rejects_too_many_rows():
    lines = ["ID"] + [f"TC_{i}" for i in range(MAX_IMPORTED_ROWS + 1)]
    uploaded = FakeUploadedFile("huge.csv", "\n".join(lines).encode("utf-8"))

    with pytest.raises(FileImportError, match="row limit"):
        parse_uploaded_file(uploaded)


def test_parse_uploaded_csv_accepts_exactly_max_rows():
    lines = ["ID"] + [f"TC_{i}" for i in range(MAX_IMPORTED_ROWS)]
    uploaded = FakeUploadedFile("boundary.csv", "\n".join(lines).encode("utf-8"))

    rows = parse_uploaded_file(uploaded)

    assert len(rows) == MAX_IMPORTED_ROWS


def test_parse_uploaded_file_rejects_unsupported_extension_without_reading_bytes():
    class ExplodingUploadedFile:
        name = "cases.txt"

        def getvalue(self):
            raise AssertionError("getvalue() must not be called for an unsupported type")

    with pytest.raises(FileImportError, match="Unsupported file type"):
        parse_uploaded_file(ExplodingUploadedFile())


def test_parse_xlsx_pads_short_data_row_with_none(monkeypatch):
    """A malformed file can yield a row shorter than the header row."""
    import core.file_import as file_import

    class _FakeSheet:
        max_column = 3

        def iter_rows(self, values_only=True):
            yield ("ID", "Title", "Extra")
            yield ("TC_001", "Login works")  # short row

    class _FakeWorkbook:
        active = _FakeSheet()

        def close(self):
            pass

    monkeypatch.setattr(
        file_import.openpyxl, "load_workbook", lambda *args, **kwargs: _FakeWorkbook()
    )

    rows = parse_uploaded_file(FakeUploadedFile("short.xlsx", b"ignored"))

    assert rows == [{"ID": "TC_001", "Title": "Login works", "Extra": None}]


def test_parse_uploaded_csv_rejects_invalid_utf8():
    uploaded = FakeUploadedFile("bad.csv", b"\x80\x81\x82")

    with pytest.raises(FileImportError, match="UTF-8"):
        parse_uploaded_file(uploaded)


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_parse_uploaded_file_rejects_duplicate_headers(extension):
    if extension == "csv":
        content = b"ID,ID\nTC_001,TC_002\n"
    else:
        content = _make_xlsx_bytes(["ID", "ID"], [["TC_001", "TC_002"]])

    with pytest.raises(FileImportError, match="duplicate column headers.*ID"):
        parse_uploaded_file(FakeUploadedFile(f"duplicate.{extension}", content))


def test_parse_uploaded_csv_rejects_row_wider_than_headers():
    uploaded = FakeUploadedFile("wide.csv", b"ID,Title\nTC_001,Login works,unexpected\n")

    with pytest.raises(FileImportError, match="more values than the header row"):
        parse_uploaded_file(uploaded)
