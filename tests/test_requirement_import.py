"""Reading requirement files and splitting them into sections (no AI involved)."""
import io
from types import SimpleNamespace

import docx
import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from core.requirement_import import (
    RequirementImportError,
    read_requirement_file,
    split_text,
)


def _file(name: str, data: bytes | str):
    raw = data.encode("utf-8") if isinstance(data, str) else data
    return SimpleNamespace(name=name, getvalue=lambda: raw)


def _docx(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf(pages: list[str], outline: list[tuple[str, int]] = ()) -> bytes:
    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    }))
    for text in pages:
        page = writer.add_blank_page(612, 792)
        shown = " ".join(f"({line}) Tj T*" for line in text.split("\n"))
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td 14 TL {shown} ET".encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
    for title, page_index in outline:
        writer.add_outline_item(title, page_index)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


# ---------- Markdown / text ----------

MARKDOWN = """# Shop requirements

Version 2.

## Login

As a user I want to log in with OTP.

AC-1: a valid OTP logs the user in.

```
# not a heading, a shell comment
```

### Lockout

Five wrong OTPs lock the account.

## Checkout

Pay by card.
"""


def test_markdown_splits_at_the_repeated_heading_level_and_keeps_subsections():
    sections = read_requirement_file(_file("spec.md", MARKDOWN))

    assert [section.title for section in sections] == ["Login", "Checkout"]
    assert [section.key for section in sections] == ["S1", "S2"]
    login = sections[0].text
    # The short title and preamble give the first section its context.
    assert login.startswith("Shop requirements\n\nVersion 2.\n\nLogin")
    assert "# not a heading, a shell comment" in login
    assert "Lockout\n\nFive wrong OTPs lock the account." in login
    assert sections[1].text == "Checkout\n\nPay by card."


def test_a_long_preamble_is_its_own_section():
    text = "Intro paragraph. " * 40 + "\n\n## A\n\nText A\n\n## B\n\nText B\n"

    sections = read_requirement_file(_file("spec.txt", text))

    assert [section.title for section in sections] == ["Introduction", "A", "B"]


def test_text_without_headings_is_one_section_or_parts_under_the_limit():
    assert [s.title for s in read_requirement_file(_file("story.txt", "As a user..."))] == ["Whole document"]

    paragraphs = "\n\n".join(f"Paragraph {n} " + "x" * 80 for n in range(30))
    sections = read_requirement_file(_file("story.txt", paragraphs), max_chars=500)

    assert len(sections) > 1
    assert all(len(section.text) <= 500 for section in sections)
    assert sections[0].title == f"Whole document (part 1 of {len(sections)})"
    assert "\n\n".join(section.text for section in sections) == paragraphs


def test_an_oversized_section_is_split_at_paragraphs():
    body = "\n\n".join("Rule " + "y" * 90 for _ in range(10))
    text = f"## Big\n\n{body}\n\n## Small\n\nOne rule."

    sections = read_requirement_file(_file("spec.md", text), max_chars=300)

    assert [s.title for s in sections][-1] == "Small"
    big = [s for s in sections if s.title.startswith("Big (part ")]
    assert len(big) >= 2 and all(len(s.text) <= 300 for s in sections)
    assert len({s.key for s in sections}) == len(sections)


def test_split_text_cuts_an_overlong_line_at_spaces_then_anywhere():
    assert split_text("word " * 50, 30) and all(len(p) <= 30 for p in split_text("word " * 50, 30))
    assert split_text("z" * 25, 10) == ["z" * 10, "z" * 10, "z" * 5]
    assert split_text("   ", 10) == []


def test_text_that_is_not_utf8_is_reported():
    with pytest.raises(RequirementImportError, match="UTF-8"):
        read_requirement_file(_file("spec.md", b"\xff\xfe\x00bad"))


@pytest.mark.parametrize("name", ["spec.xlsx", "spec.doc", "spec"])
def test_unsupported_types_are_rejected(name):
    with pytest.raises(RequirementImportError, match="Unsupported file type"):
        read_requirement_file(_file(name, "text"))


def test_files_over_the_size_limit_are_rejected():
    with pytest.raises(RequirementImportError, match="10 MiB"):
        read_requirement_file(_file("spec.txt", b"a" * (10 * 1024 * 1024 + 1)))


def test_an_empty_file_has_no_sections():
    with pytest.raises(RequirementImportError, match="No text"):
        read_requirement_file(_file("spec.md", "\n\n  \n"))


# ---------- Word ----------

def test_docx_splits_at_heading_styles_and_keeps_tables_and_lists():
    def build(document):
        document.add_heading("Payment spec", level=0)
        document.add_heading("Card payment", level=1)
        document.add_paragraph("The user pays by card.")
        document.add_paragraph("Card number is validated", style="List Number")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Field", "Rule"
        table.cell(1, 0).text, table.cell(1, 1).text = "CVV", "3 digits"
        document.add_heading("Refunds", level=1)
        document.add_paragraph("Refunds within 30 days.")

    sections = read_requirement_file(_file("spec.docx", _docx(build)))

    assert [s.title for s in sections] == ["Card payment", "Refunds"]
    card = sections[0].text
    assert card.startswith("Payment spec")
    assert "- Card number is validated" in card
    assert "Field | Rule\nCVV | 3 digits" in card
    assert sections[1].text == "Refunds\n\nRefunds within 30 days."


def test_a_broken_docx_is_reported():
    with pytest.raises(RequirementImportError, match="Word file"):
        read_requirement_file(_file("spec.docx", b"not a zip"))


# ---------- PDF ----------

def test_pdf_with_bookmarks_splits_at_each_bookmark():
    data = _pdf(
        ["Cover page", "Login\nOTP logs the user in", "Lockout after 5 tries", "Checkout\nPay by card"],
        outline=[("Login", 1), ("Checkout", 3)],
    )

    sections = read_requirement_file(_file("spec.pdf", data))

    assert [s.title for s in sections] == ["Login", "Checkout"]
    assert "Cover page" in sections[0].text
    assert "OTP logs the user in" in sections[0].text and "Lockout after 5 tries" in sections[0].text
    assert "Pay by card" in sections[1].text


def test_pdf_without_bookmarks_is_one_section():
    sections = read_requirement_file(_file("spec.pdf", _pdf(["First page", "Second page"])))

    assert len(sections) == 1
    assert "First page" in sections[0].text and "Second page" in sections[0].text


def test_a_pdf_without_text_or_a_broken_pdf_is_reported():
    with pytest.raises(RequirementImportError, match="scanned PDF"):
        read_requirement_file(_file("scan.pdf", _pdf([""])))
    with pytest.raises(RequirementImportError, match="PDF"):
        read_requirement_file(_file("spec.pdf", b"%PDF-1.4 garbage"))


# ---------- CSV backlog ----------

def test_csv_backlog_is_one_section_per_story():
    csv_text = (
        "Issue key,Summary,Description,Acceptance Criteria\n"
        "US-1,Log in with OTP,As a user...,\"AC-1 valid OTP\nAC-2 wrong OTP\"\n"
        ",,,\n"
        "US-2,Checkout,As a buyer...,AC-1 pay by card\n"
        "US-2,Refund,As a buyer...,AC-1 refund\n"
    )

    sections = read_requirement_file(_file("backlog.csv", csv_text))

    assert [s.key for s in sections] == ["US-1", "US-2", "US-2_2"]
    assert sections[0].title == "Row 2: US-1 Log in with OTP"
    assert sections[0].text == (
        "Issue key: US-1\nSummary: Log in with OTP\nDescription: As a user...\n"
        "Acceptance Criteria: AC-1 valid OTP\nAC-2 wrong OTP"
    )
    assert sections[2].title.startswith("Row 6: US-2 Refund")


def test_csv_without_id_columns_uses_row_numbers_and_s_keys():
    sections = read_requirement_file(_file("backlog.csv", "Story\nAs a user I log in\nAs a user I log out\n"))

    assert [s.key for s in sections] == ["S1", "S2"]
    assert sections[0].title == "Row 2: As a user I log in"


def test_csv_errors_come_from_the_shared_parser():
    with pytest.raises(RequirementImportError, match="no data rows"):
        read_requirement_file(_file("backlog.csv", "Story\n"))
