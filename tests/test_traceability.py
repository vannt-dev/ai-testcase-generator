"""Traceability: criteria -> test cases, in core, the Excel sheet and the Gherkin tags."""
import copy
import io

import pytest
from openpyxl import load_workbook

from core.excel_exporter import export_to_excel
from core.gherkin_exporter import export_to_gherkin
from core.traceability import build_traceability, has_traceability, uncovered_criteria


def _case(test_id, covers, **overrides):
    case = {
        "test_id": test_id,
        "module": "Login",
        "title": f"Case {test_id}",
        "precondition": "Account exists",
        "steps": "1. Log in",
        "test_data": "Valid account",
        "expected_result": "Home page opens",
        "priority": "High",
        "type": "Positive",
        "platform": "Web",
        "covers": covers,
    }
    case.update(overrides)
    return case


RESULT = {
    "acceptance_criteria": [
        {"id": "AC-1", "text": "Log in with phone and OTP"},
        {"id": "AC-2", "text": "Wrong OTP is rejected"},
        {"id": "AC-3", "text": "OTP expires after 5 minutes"},
    ],
    "test_cases": [
        _case("TC_001", ["AC-1"]),
        _case("TC_002", ["AC-2", "AC-1"]),
        _case("TC_003", "AC-9"),
        _case("TC_004", []),
    ],
    "summary": {"total": 4, "by_type": {"positive": 4}, "open_questions": []},
}


def test_build_traceability_maps_each_criterion_to_its_test_cases():
    trace = build_traceability(RESULT)

    assert [(row["id"], row["test_ids"], row["listed"]) for row in trace["criteria"]] == [
        ("AC-1", ["TC_001", "TC_002"], True),
        ("AC-2", ["TC_002"], True),
        ("AC-3", [], True),
        ("AC-9", ["TC_003"], False),
    ]
    assert trace["criteria"][0]["text"] == "Log in with phone and OTP"
    assert trace["unlinked_test_ids"] == ["TC_004"]
    assert uncovered_criteria(RESULT) == ["AC-3"]


def test_results_without_links_have_no_traceability():
    legacy = {"test_cases": [{k: v for k, v in _case("TC_1", []).items() if k != "covers"}]}

    assert not has_traceability(legacy)
    assert build_traceability(legacy) == {"criteria": [], "unlinked_test_ids": ["TC_1"]}
    assert has_traceability(RESULT)


def test_duplicate_and_blank_criterion_ids_are_listed_once():
    result = {
        "acceptance_criteria": [{"id": "AC-1", "text": "a"}, {"id": " AC-1 ", "text": "b"}, {"id": "", "text": "c"}],
        "test_cases": [],
    }

    assert [row["id"] for row in build_traceability(result)["criteria"]] == ["AC-1"]


def test_excel_adds_a_covers_column_and_a_traceability_sheet_that_flags_gaps():
    workbook = load_workbook(io.BytesIO(export_to_excel(RESULT)))

    cases = workbook["Test Cases"]
    assert cases.cell(row=1, column=11).value == "Covers"
    assert cases.cell(row=3, column=11).value == "AC-2, AC-1"
    assert cases.cell(row=5, column=11).value is None

    sheet = workbook["Traceability"]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert rows[0] == ["Criterion", "Description", "Test IDs", "Test count", "Status"]
    assert rows[1] == ["AC-1", "Log in with phone and OTP", "TC_001, TC_002", 2, "Covered"]
    assert rows[3] == ["AC-3", "OTP expires after 5 minutes", None, 0, "No test case"]
    assert rows[4][0] == "AC-9" and rows[4][4] == "Covered (not in the criteria list)"
    assert ["Test cases not linked to a criterion", None, None, None, None] in rows
    assert ["TC_004", None, None, None, None] in rows
    assert sheet.cell(row=4, column=1).fill.start_color.rgb == "00FFC7CE"
    assert sheet.cell(row=2, column=1).fill.start_color.rgb != "00FFC7CE"


def test_excel_without_links_keeps_the_original_layout():
    legacy = copy.deepcopy(RESULT)
    del legacy["acceptance_criteria"]
    for case in legacy["test_cases"]:
        del case["covers"]

    workbook = load_workbook(io.BytesIO(export_to_excel(legacy)))

    assert "Traceability" not in workbook.sheetnames
    assert workbook["Test Cases"].max_column == 10


@pytest.mark.parametrize("text", ["=1+1", "@AC"])
def test_traceability_sheet_keeps_text_literal(text):
    result = {
        "acceptance_criteria": [{"id": text, "text": text}],
        "test_cases": [_case(text, [text])],
        "summary": {},
    }

    sheet = load_workbook(io.BytesIO(export_to_excel(result)))["Traceability"]

    assert sheet["A2"].value == text and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == text and sheet["B2"].data_type == "s"
    assert sheet["C2"].value == text and sheet["C2"].data_type == "s"


def test_gherkin_tags_each_scenario_with_the_criteria_it_covers():
    feature = export_to_gherkin(RESULT, feature_name="Login")

    assert "  @TC_001 @Login @High @Positive @Web @AC-1\n" in feature
    assert "  @TC_002 @Login @High @Positive @Web @AC-2 @AC-1\n" in feature
    assert "  @TC_004 @Login @High @Positive @Web\n" in feature


def test_gherkin_tag_from_an_id_with_spaces_stays_one_token():
    result = {"test_cases": [_case("TC_1", ["REQ 1.2"])]}

    assert "@REQ_1.2" in export_to_gherkin(result)
